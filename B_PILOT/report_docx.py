"""Render the report to a ``.docx`` with its figures embedded.

Built for the Google Docs backend, and it is what makes figures work there at
all. Drive converts an uploaded ``.docx`` into a native Google Doc, and images
inside a ``.docx`` are *real files in the archive* rather than links or base64
in markup -- so they survive the conversion as inline pictures. Uploading the
Markdown instead is simpler but drops every figure on the floor, because a
relative ``figures/x.png`` path means nothing to Drive.

**Why a third renderer, when to_html and render_markdown already exist.**
Because this one has a different contract, and the two existing ones each fail
it for a structural reason rather than an incidental one:

* ``report_render.to_html`` draws for the *screen*: every colour comes from the
  session's theme, so a dark session would produce a near-white document. The
  established fix (``style.temporary_theme``, used by the PDF exporter) rebinds
  module globals, which is safe on the GUI thread and emphatically not safe on
  the sync worker's thread while the GUI is painting.
* ``report_builder.render_markdown`` is the source text, not a document format.

So this module maps the Markdown to Word's own semantics -- headings are
headings, tables are tables, figures are pictures -- and carries no palette at
all. That also makes it Qt-free and unit-testable without a display.

The grammar handled is exactly what :func:`report_builder.render_markdown`
emits, and is kept deliberately in step with ``report_render.to_html``'s: HTML
comments (dropped), ATX headings, ``---`` rules, fenced code, pipe tables,
``>`` quotes, ``![](figures/...)`` figures, and paragraphs with inline
code/bold/italic. Anything unrecognised falls through as a paragraph rather
than being lost -- a lab record must never silently drop a line.

Links are rendered as **inert text**, not hyperlinks, matching
``report_render._link_html``. The body of this document is arbitrary text from
the kernel, from a pasted note, or drafted by AutoPILOT; the report's own
``bpilot:`` control links are stripped, and nothing else is allowed to become
clickable in a document that gets forwarded to other people.
"""
from __future__ import annotations

import io
import os
import re

MISSING_REASON = ""
try:
    import docx as _docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH as _ALIGN
    from docx.shared import Inches as _Inches
    from docx.shared import Pt as _Pt
    from docx.shared import RGBColor as _RGBColor
except Exception as exc:  # noqa: BLE001 -- absent, or a broken partial install
    MISSING_REASON = f"{type(exc).__name__}: {exc}"

# Kept in step with report_render's, deliberately: the two renderers consume
# the same emitted grammar and drifting apart would mean one of them silently
# mishandling a construct the other supports.
_CODE = re.compile(r"`([^`]+)`")
_BOLD = re.compile(r"\*\*([^*]+)\*\*")
_ITALIC = re.compile(r"(?<![\w*])_([^_]+)_(?![\w*])")
_LINK = re.compile(r"(?<!!)\[([^\]]+)\]\(([^)\s]+)\)")
_IMAGE = re.compile(r'^!\[([^\]]*)\]\(\s*([^)\s]+)(?:\s+"[^"]*")?\s*\)\s*$')
_FENCE = re.compile(r"^```\s*([A-Za-z0-9_+-]*)\s*$")
_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
_TABLE_SEP = re.compile(r"^\|[\s:|-]+\|$")
_COMMENT_OPEN = re.compile(r"^\s*<!--")
_COMMENT_CLOSE = re.compile(r"-->")

#: Widest a figure is placed, in inches. A US-Letter page with one-inch
#: margins is 6.5in of text, so this fills the column without overflowing it.
FIGURE_WIDTH_IN = 6.0

_MUTED = (0x60, 0x60, 0x60)


def available() -> bool:
    return not MISSING_REASON


def _cells(line: str) -> list:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def _strip_controls(text: str) -> str:
    """Drop the report's own ``bpilot:`` control links.

    They are chrome for the live panel (hide, move, unhide) and meaningless in
    an exported document -- ``to_html`` only emits them when ``controls=True``,
    which no export path sets, but stripping here keeps this renderer correct
    even if it is ever handed a document that has them.
    """
    return _LINK.sub(lambda m: "" if m.group(2).startswith("bpilot:") else m.group(0), text)


def _add_runs(paragraph, text: str) -> None:
    """Append `text` to `paragraph`, honouring the inline spans we emit.

    One pass, left to right, so a span cannot be applied twice and the
    remaining literal text between matches is preserved exactly.
    """
    text = _strip_controls(text)
    # Links become their label plus the target in parentheses -- inert, and it
    # still tells the reader where the author was pointing.
    text = _LINK.sub(lambda m: f"{m.group(1)} ({m.group(2)})", text)

    pattern = re.compile(
        f"({_CODE.pattern})|({_BOLD.pattern})|({_ITALIC.pattern})"
    )
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos:match.start()])
        body = match.group(2) or match.group(4) or match.group(6) or ""
        run = paragraph.add_run(body)
        if match.group(1):
            run.font.name = "Courier New"
        elif match.group(3):
            run.bold = True
        else:
            run.italic = True
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def _add_code(document, lines: list) -> None:
    para = document.add_paragraph()
    para.paragraph_format.left_indent = _Inches(0.25)
    para.paragraph_format.space_after = _Pt(6)
    run = para.add_run("\n".join(lines))
    run.font.name = "Courier New"
    run.font.size = _Pt(9)


def _add_quote(document, lines: list) -> None:
    para = document.add_paragraph()
    para.paragraph_format.left_indent = _Inches(0.35)
    run = para.add_run(" ".join(line.lstrip("> ").rstrip() for line in lines))
    run.italic = True
    run.font.color.rgb = _RGBColor(*_MUTED)


def _add_rule(document) -> None:
    para = document.add_paragraph()
    para.alignment = _ALIGN.CENTER
    run = para.add_run("• • •")
    run.font.color.rgb = _RGBColor(*_MUTED)


def _add_table(document, rows: list) -> None:
    if not rows:
        return
    width = max(len(r) for r in rows)
    table = document.add_table(rows=0, cols=width)
    table.style = "Table Grid"
    for index, row in enumerate(rows):
        cells = table.add_row().cells
        for column in range(width):
            text = row[column] if column < len(row) else ""
            para = cells[column].paragraphs[0]
            _add_runs(para, text)
            if index == 0:
                for run in para.runs:
                    run.bold = True


def _add_figure(document, alt: str, rel: str, base_dir: str) -> None:
    """Place one figure, or say plainly that it is missing.

    A record that lost a file should say so rather than showing a gap -- the
    same promise ``report_render._image_html`` makes on screen.
    """
    path = os.path.join(base_dir, rel) if base_dir else rel
    if not os.path.isfile(path):
        para = document.add_paragraph()
        run = para.add_run(f"[missing figure: {rel}]")
        run.italic = True
        run.font.color.rgb = _RGBColor(*_MUTED)
        return
    try:
        para = document.add_paragraph()
        para.alignment = _ALIGN.CENTER
        para.add_run().add_picture(path, width=_Inches(FIGURE_WIDTH_IN))
    except Exception:  # noqa: BLE001 -- unreadable, or a format Word rejects
        para = document.add_paragraph()
        run = para.add_run(f"[unreadable figure: {rel}]")
        run.italic = True
        run.font.color.rgb = _RGBColor(*_MUTED)
        return
    if alt:
        caption = document.add_paragraph()
        caption.alignment = _ALIGN.CENTER
        run = caption.add_run(alt)
        run.italic = True
        run.font.size = _Pt(9)
        run.font.color.rgb = _RGBColor(*_MUTED)


def build(markdown: str, *, base_dir: str = "", title: str = "") -> bytes:
    """Render `markdown` to ``.docx`` bytes, embedding figures from `base_dir`.

    Raises nothing the caller has to catch beyond the obvious: a sink wraps
    this, because a malformed record must become a retryable failure rather
    than kill the worker thread.
    """
    document = _docx.Document()
    if title:
        document.core_properties.title = title

    lines = (markdown or "").splitlines()
    i = 0
    in_comment = False

    while i < len(lines):
        line = lines[i]

        if in_comment:
            in_comment = not _COMMENT_CLOSE.search(line)
            i += 1
            continue
        if _COMMENT_OPEN.match(line):
            in_comment = not _COMMENT_CLOSE.search(line)
            i += 1
            continue

        if not line.strip():
            i += 1
            continue

        if _FENCE.match(line.strip()):
            i += 1
            block: list = []
            while i < len(lines) and not _FENCE.match(lines[i].strip()):
                block.append(lines[i])
                i += 1
            i += 1  # closing fence, or EOF -- both fine
            _add_code(document, block)
            continue

        if line.strip() in ("---", "***", "___"):
            _add_rule(document)
            i += 1
            continue

        heading = _HEADING.match(line)
        if heading:
            text = _strip_controls(heading.group(2)).strip()
            if text:
                document.add_heading(text, level=min(len(heading.group(1)), 4))
            i += 1
            continue

        image = _IMAGE.match(line.strip())
        if image:
            _add_figure(document, image.group(1), image.group(2), base_dir)
            i += 1
            continue

        if line.lstrip().startswith(">"):
            block = []
            while i < len(lines) and lines[i].lstrip().startswith(">"):
                block.append(lines[i])
                i += 1
            _add_quote(document, block)
            continue

        if line.lstrip().startswith("|"):
            rows = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                if not _TABLE_SEP.match(lines[i].strip()):
                    rows.append(_cells(lines[i]))
                i += 1
            _add_table(document, rows)
            continue

        _add_runs(document.add_paragraph(), line)
        i += 1

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
